#!/usr/bin/env python3
"""Export PRD markdown to a styled Word document."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HEADING_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
BODY_COLOR = RGBColor(0x33, 0x33, 0x33)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, *, bold=False, italic=False, size_pt=11, color=None, name=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color
    if name:
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline_markdown(paragraph, text: str, *, base_size=11, base_color=BODY_COLOR):
    """Parse **bold**, *italic*, and `code` into runs."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, size_pt=base_size, color=base_color)
        chunk = m.group(0)
        if chunk.startswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, bold=True, size_pt=base_size, color=base_color)
        elif chunk.startswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, size_pt=10, color=base_color, name="Menlo")
        else:
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, italic=True, size_pt=base_size, color=base_color)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size_pt=base_size, color=base_color)


def style_paragraph(paragraph, *, space_before=0, space_after=6, line_spacing=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing


def add_heading(doc: Document, text: str, level: int):
    clean = re.sub(r"^#+\s*", "", text).strip()
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}.get(level, "Heading 4")
    p = doc.add_paragraph(style=style)
    p.clear()
    run = p.add_run(clean)
    sizes = {1: 22, 2: 16, 3: 13, 4: 12}
    set_run_font(run, bold=True, size_pt=sizes.get(level, 12), color=HEADING_COLOR)
    spaces = {1: (0, 12), 2: (18, 8), 3: (14, 6), 4: (10, 4)}
    sb, sa = spaces.get(level, (10, 4))
    style_paragraph(p, space_before=sb, space_after=sa)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\|[\s\-:|]+\|$", line.strip()))


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if is_separator_row(line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.rows[ri].cells[ci]
            text = row[ci] if ci < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.clear()
            if ri == 0:
                run = p.add_run(re.sub(r"\*\*", "", text))
                set_run_font(run, bold=True, size_pt=10, color=HEADING_COLOR)
            else:
                add_inline_markdown(p, text, base_size=10)
    doc.add_paragraph()


def add_bullet(doc: Document, text: str, *, ordered: bool = False, level: int = 0):
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.clear()
    add_inline_markdown(p, text.lstrip(), base_size=11)
    style_paragraph(p, space_before=0, space_after=3)
    p.paragraph_format.left_indent = Inches(0.25 * level)


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    add_inline_markdown(p, text)
    style_paragraph(p)


def add_hr(doc: Document):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=6, space_after=6)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        from docx.oxml import OxmlElement

        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        from docx.oxml import OxmlElement

        bottom = OxmlElement("w:bottom")
        pBdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_COLOR
    for name, size in [("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 13), ("Heading 4", 12)]:
        if name in doc.styles:
            st = doc.styles[name]
            st.font.name = "Calibri"
            st.font.bold = True
            st.font.size = Pt(size)
            st.font.color.rgb = HEADING_COLOR


def markdown_to_docx(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            add_hr(doc)
            i += 1
            continue

        if is_table_row(stripped):
            block = []
            while i < len(lines) and is_table_row(lines[i].strip()):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table_rows(block))
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            if level == 1 and not title_done:
                add_heading(doc, stripped, 1)
                title_done = True
            elif stripped.startswith("## "):
                add_heading(doc, stripped, 2)
            elif stripped.startswith("### "):
                add_heading(doc, stripped, 3)
            elif stripped.startswith("#### "):
                add_heading(doc, stripped, 4)
            else:
                add_heading(doc, stripped, min(level, 4))
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            add_bullet(doc, text)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            add_bullet(doc, text, ordered=True)
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            set_run_font(run, size_pt=9, name="Menlo", color=MUTED_COLOR)
            style_paragraph(p, space_before=4, space_after=8)
            continue

        add_body(doc, stripped)
        i += 1

    doc.save(out_path)


def main():
    root = Path(__file__).resolve().parents[1]
    md = root / "docs" / "PRD.md"
    out = Path.home() / "Desktop" / "Car-Comfort-PRD.docx"
    if len(sys.argv) > 1:
        out = Path(sys.argv[1])
    markdown_to_docx(md, out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
